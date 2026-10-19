"""Fill the official ADCAIJ Word template (ADCAIJ_Template_2026_V6.dotx) with the manuscript content.
Usage: python fill_template.py <template.dotx> <blocks.json> <source_manuscript.docx (for the OMML equation)> <out.docx>
blocks.json is written by build_manuscript.js (semantic record of every heading, paragraph, table and figure).
"""
import copy, json, re, sys, shutil, tempfile, zipfile, os
from docx import Document
from docx.shared import Pt, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

tpl, blocks_path, src_docx, out = sys.argv[1:5]
D = json.load(open(blocks_path))
BLIND = D['blind']

# 1. .dotx -> .docx (only the main content type differs)
tmp = tempfile.mkdtemp()
docx_in = os.path.join(tmp, 'tpl.docx')
with zipfile.ZipFile(tpl) as zi, zipfile.ZipFile(docx_in, 'w', zipfile.ZIP_DEFLATED) as zo:
    for it in zi.infolist():
        data = zi.read(it.filename)
        if it.filename == '[Content_Types].xml':
            data = data.replace(b'wordprocessingml.template.main+xml', b'wordprocessingml.document.main+xml')
        zo.writestr(it, data)
doc = Document(docx_in)
body = doc.element.body

def style_of(p):
    return p.style.style_id if p.style is not None else None

def add_runs(p, s, size=None, italic_base=False):
    for part in re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)', s):
        if not part:
            continue
        if part.startswith('**'):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith('*'):
            r = p.add_run(part[1:-1]); r.italic = not italic_base
        else:
            r = p.add_run(part)
            if italic_base: r.italic = True
        if size: r.font.size = Pt(size)
    return p

def set_text(p, s):
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    for h in p._p.findall(qn('w:hyperlink')):
        p._p.remove(h)
    return add_runs(p, s)

# 2. front matter (some front-matter paragraphs sit inside tables, so walk the whole body)
from docx.text.paragraph import Paragraph as _P
def all_paras():
    return [_P(e, doc._body) for e in body.iter(qn('w:p'))]
paras = all_paras()
AUTH = [('Hilman Firdaus', '0009-0006-8820-3305', 'hilmanfirdaus48@gmail.com'),
        ('Wendy Sarasjati', '0009-0009-8679-0050', 'wendysar@unimus.ac.id'),
        ('Dhendra Marutho', '0009-0004-2271-1767', 'dhendra@unimus.ac.id'),
        ('Ahmad Ilham', '0000-0001-5109-6258', 'ahmadilham@unimus.ac.id')]
AFF = ('Magister Programme of Informatics, Universitas Muhammadiyah Semarang (Gedung NRC), '
       'Jl. Kedungmundu Raya No. 18, Semarang 50273, Central Java, Indonesia')
title_p = next(p for p in paras if style_of(p) == 'Ttulo10')
set_text(title_p, D['title'])
auth_p = next(p for p in paras if style_of(p) == 'AuthorsName')
set_text(auth_p, '')
if BLIND:
    auth_p.add_run('[Authors removed for double-blind review]')
else:
    for i, (nm, orcid, _) in enumerate(AUTH):
        auth_p.add_run(nm)
        sup = auth_p.add_run('a,*' if i == 0 else 'a'); sup.font.superscript = True
        auth_p.add_run(f' [{orcid}]' + (', ' if i < len(AUTH) - 2 else (' and ' if i == len(AUTH) - 2 else '')))
affs = [p for p in paras if style_of(p) == 'affiliations']
set_text(affs[0], '')
if BLIND:
    affs[0].add_run('[Affiliations removed for double-blind review]')
else:
    s = affs[0].add_run('a'); s.font.superscript = True
    affs[0].add_run(' ' + AFF)
for p in affs[1:]:
    p._p.getparent().remove(p._p)
em = next(p for p in paras if style_of(p) == 'emails')
if BLIND:
    set_text(em, '[E-mail addresses removed for double-blind review]')
else:
    set_text(em, ', '.join(a[2] for a in AUTH))
    cp = copy.deepcopy(em._p); em._p.addnext(cp)
    from docx.text.paragraph import Paragraph
    cpp = Paragraph(cp, em._parent)
    set_text(cpp, '* Corresponding author: Hilman Firdaus, hilmanfirdaus48@gmail.com. Postal address as in affiliation a. '
                  'Tel. +62 24 76740287.')
kws = [p for p in all_paras() if style_of(p) == 'Keywords']
set_text(kws[0], '; '.join(D['keywords'])); kws[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
for p in kws[1:]:
    p._p.getparent().remove(p._p)
absp = [p for p in all_paras() if style_of(p) == 'Abstract']
target = next((p for p in absp if p.text.strip()), absp[-1])
set_text(target, D['abstract'])
for p in absp:
    if p is not target and not p.text.strip():
        p._p.getparent().remove(p._p)

# 3. remove template body after the keyword/abstract table
tbl = [t for t in body.findall(qn('w:tbl')) if 'Keywords' in ''.join(e.get(qn('w:val')) or '' for e in t.iter(qn('w:pStyle')))][0]
el = tbl.getnext()
while el is not None and el.tag != qn('w:sectPr'):
    nxt = el.getnext(); body.remove(el); el = nxt
sectPr = body.find(qn('w:sectPr'))

sec = doc.sections[0]
text_width = sec.page_width - sec.left_margin - sec.right_margin

def new_par(style=None):
    p = doc.add_paragraph()
    if style: p.style = doc.styles.get_by_id(style, 1) if False else None
    return p

from docx.enum.style import WD_STYLE_TYPE
SID = {s.style_id: s for s in doc.styles if s.type == WD_STYLE_TYPE.PARAGRAPH}
def para(style_id=None, text=None, **kw):
    p = doc.add_paragraph()
    if style_id: p.style = SID[style_id]
    if text is not None: add_runs(p, text, **kw)
    return p

def border(cell, edges):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement('w:tcBorders')
    for e in ('top', 'left', 'bottom', 'right'):
        x = OxmlElement(f'w:{e}')
        if e in edges:
            x.set(qn('w:val'), 'single'); x.set(qn('w:sz'), '6'); x.set(qn('w:color'), '000000')
        else:
            x.set(qn('w:val'), 'nil')
        b.append(x)
    tcPr.append(b)

src = Document(src_docx)
eq_p = next(p for p in src.paragraphs if p._p.find('.//' + qn('m:oMath')) is not None)

first = True
for b in D['blocks']:
    t = b['t']
    if t == 'h1':
        txt = re.sub(r'^\d+\.\s*', '', b['text'])
        hp = para('Ttulo3', txt); first = True
        if txt == 'References': hp.paragraph_format.page_break_before = True  # template: references start on a new page
    elif t == 'h2':
        para('Ttulo4', re.sub(r'^\d+\.\d+\.\s*', '', b['text'])); first = True
    elif t == 'p':
        s = b['s']
        if s.startswith('**Acknowledgments.**'):
            para('Ttulo3', 'Acknowledgements')
            para('p1a', '[Removed for double-blind review]' if BLIND else s.replace('**Acknowledgments.** ', ''))
            first = False; continue
        para('p1a' if first else None, s); first = False
    elif t == 'ol':
        for it in b['items']:
            para('Prrafodelista', it)
        first = False
    elif t == 'eq':
        doc.add_paragraph()._p.addprevious(copy.deepcopy(eq_p._p))
        body.remove(body[-2]) if False else None
        # drop the empty helper paragraph just added
        last = [e for e in body if e.tag == qn('w:p')][-1]; body.remove(last)
        first = False
    elif t == 'table':
        cap = para('Descripcin', f"Table {b['n']}. {b['caption']}"); cap.paragraph_format.keep_with_next = True
        ncol = len(b['header'])
        tb = doc.add_table(rows=1 + len(b['rows']), cols=ncol)
        tb.alignment = WD_TABLE_ALIGNMENT.CENTER
        tot = sum(b['widths'])
        tb.autofit = False
        tw_dxa = int(text_width / 635)  # EMU -> twentieths of a point
        cols = [int(tw_dxa * w / tot) for w in b['widths']]; cols[-1] += tw_dxa - sum(cols)
        tblPr = tb._tbl.tblPr
        tblW = tblPr.find(qn('w:tblW'))
        if tblW is None: tblW = OxmlElement('w:tblW'); tblPr.append(tblW)
        tblW.set(qn('w:w'), str(tw_dxa)); tblW.set(qn('w:type'), 'dxa')
        lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
        grid = tb._tbl.tblGrid
        for gc, w in zip(grid.findall(qn('w:gridCol')), cols): gc.set(qn('w:w'), str(w))
        for row in tb.rows:
            for j, c in enumerate(row.cells): c.width = Emu(cols[j] * 635)
        tb.rows[0].repeat_header = True if hasattr(tb.rows[0], 'repeat_header') else None
        for i, row in enumerate([b['header']] + b['rows']):
            last = i == len(b['rows'])
            for j, val in enumerate(row):
                c = tb.rows[i].cells[j]; c.text = ''
                p = c.paragraphs[0]; add_runs(p, val, size=9)
                if i == 0:
                    for r in p.runs: r.bold = True
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(0); p.paragraph_format.first_line_indent = Pt(0)
                border(c, {'top', 'bottom'} if i == 0 else ({'bottom'} if last else set()))
        nt = para(None, b['note'], size=9); nt.paragraph_format.first_line_indent = Pt(0); first = True
    elif t == 'figure':
        p = para(None); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(b['file'], width=Emu(int(text_width * min(1.0, b['width'] / 600))))
        para('Descripcin', f"Figure {b['n']}. {b['caption']}")
        nt = para(None, 'Source: Own elaboration.', size=9); nt.alignment = WD_ALIGN_PARAGRAPH.CENTER; nt.paragraph_format.first_line_indent = Pt(0); first = True
    elif t == 'refs':
        for r in b['refs']:
            para('references', r)
# keep sectPr last
body.remove(sectPr); body.append(sectPr)

# 4. running header and document properties
for s in doc.sections:
    for hdr in (s.header, s.first_page_header, s.even_page_header):
        for p in hdr.paragraphs:
            for r in p.runs:
                if 'Author One et al.' in r.text:
                    r.text = r.text.replace('Author One et al.', 'Anonymous' if BLIND else 'Firdaus et al.')
                if 'ADCAIJ Submission Instructions Guidelines' in r.text:
                    r.text = r.text.replace('ADCAIJ Submission Instructions Guidelines', 'How Stable Are Cluster Explanations?')
cp_ = doc.core_properties
cp_.author = 'Anonymous' if BLIND else 'Hilman Firdaus'; cp_.last_modified_by = cp_.author
cp_.title = D['title']; cp_.comments = ''; cp_.subject = ''
doc.save(out)
# running header: replace placeholder text in every header part (text may be split across runs)
tmp2 = out + '.tmp'
with zipfile.ZipFile(out) as zi, zipfile.ZipFile(tmp2, 'w', zipfile.ZIP_DEFLATED) as zo:
    for it in zi.infolist():
        data = zi.read(it.filename)
        if re.match(r'word/header\d*\.xml', it.filename):
            x = data.decode('utf8')
            for old, new in (('Author One et al.', 'Anonymous' if BLIND else 'Firdaus et al.'),
                             ('ADCAIJ Submission Instructions Guidelines', 'How Stable Are Cluster Explanations?')):
                if old not in x:
                    # merge split runs: collapse text across <w:t> within each paragraph
                    def fix(m):
                        para_x = m.group(0); txt = ''.join(re.findall(r'<w:t[^>]*>([^<]*)</w:t>', para_x))
                        if old not in txt: return para_x
                        ts = list(re.finditer(r'(<w:t[^>]*>)([^<]*)(</w:t>)', para_x))
                        full = txt.replace(old, new); out_x = para_x
                        for k, t in enumerate(reversed(ts)):
                            idx = len(ts) - 1 - k
                            out_x = out_x[:t.start()] + t.group(1).replace('<w:t>', '<w:t xml:space="preserve">') + (full if idx == 0 else '') + t.group(3) + out_x[t.end():]
                        return out_x
                    x = re.sub(r'<w:p[ >].*?</w:p>', fix, x, flags=re.S)
                else:
                    x = x.replace(old, new)
            data = x.encode('utf8')
        zo.writestr(it, data)
shutil.move(tmp2, out)
print('written', out)
